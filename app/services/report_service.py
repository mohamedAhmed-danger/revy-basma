import csv
import io
import urllib.parse
from datetime import timezone
from zoneinfo import ZoneInfo

from flask import Response
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models import User, Attendance
from app.i18n import t
from app.services.attendance_service import get_employee_attendance_summary

CAIRO_TZ = ZoneInfo("Africa/Cairo")


def _to_cairo(value, fmt='%Y-%m-%d %I:%M %p'):
    """Same display-only UTC -> Africa/Cairo rule as the cairo_time Jinja filter,
    used here for the export functions instead of a template."""
    if value is None:
        return ''
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(CAIRO_TZ).strftime(fmt)


def _sanitize_cell(value):
    """SEC-19: Prevent CSV and Excel formula injection (CSV/formula injection).
    If a cell starts with =, +, -, @, \\t, or \\r, prefix with a single quote
    so spreadsheet software renders it as text rather than executing it as an expression.
    """
    if value is None:
        return ''
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value
    s = str(value)
    if s and s[0] in ('=', '+', '-', '@', '\t', '\r'):
        return "'" + s
    return s


def get_report(date_from, date_to, zone_id=None, status=None, employee_id=None):
    query = (
        Attendance.query
        .join(User, Attendance.user_id == User.id)
        .filter(User.is_active == True, User.role == 'employee')
        .filter(Attendance.date >= date_from, Attendance.date <= date_to)
    )

    if zone_id:
        query = query.filter(Attendance.zone_id == zone_id)

    if status == 'full_shift':
        query = query.filter(Attendance.shift_status == 'full_shift')
    elif status == 'short_shift':
        query = query.filter(Attendance.shift_status == 'short_shift')
    elif status == 'active':
        query = query.filter(Attendance.status == 'active')

    if employee_id:
        query = query.filter(Attendance.user_id == employee_id)

    return query.order_by(Attendance.date.desc(), Attendance.sign_in_time.desc()).all()


def _get_status_text(shift_status, status_attr, lang='ar'):
    if shift_status == 'full_shift':
        return t('export_status_full', lang=lang)
    elif shift_status == 'short_shift':
        return t('export_status_short', lang=lang)
    elif status_attr == 'active':
        return t('export_status_active', lang=lang)
    return ''


def export_csv(rows, lang='ar', date_from=None, date_to=None):
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    
    headers = [
        t('export_col_date', lang=lang),
        t('export_col_employee', lang=lang),
        t('export_col_zone', lang=lang),
        t('export_col_signin', lang=lang),
        t('export_col_signout', lang=lang),
        t('export_col_total_hours', lang=lang),
        t('export_col_shift_status', lang=lang),
        t('export_col_late', lang=lang),
        t('export_col_late_minutes', lang=lang),
        t('export_col_extra_days', lang=lang),
        t('export_col_absent_days', lang=lang),
        t('export_col_attendance_rate', lang=lang),
    ]
    writer.writerow(headers)

    yes_text = t('export_yes', lang=lang)
    no_text = t('export_no', lang=lang)
    start_d = date_from or (min(r.date for r in rows) if rows else date.today())
    end_d = date_to or (max(r.date for r in rows) if rows else date.today())
    user_summaries = {}

    for row in rows:
        if row.user_id not in user_summaries:
            user_summaries[row.user_id] = get_employee_attendance_summary(row.user, start_d, end_d)
        summary = user_summaries[row.user_id]

        is_day_off = not row.user.is_workday(row.date)
        day_off_text = yes_text if is_day_off else no_text

        absent_count = len(summary['absent_days'])
        absent_dates_str = ', '.join(d.strftime('%Y-%m-%d') for d in summary['absent_days'])
        absent_text = f"{absent_count} ({absent_dates_str})" if absent_count > 0 else t('no_absent_days', lang=lang)
        rate_text = f"{summary['attendance_rate']}%"

        writer.writerow([
            row.date.isoformat(),
            _sanitize_cell(row.user.name),
            _sanitize_cell(row.zone.name if row.zone else ''),
            _to_cairo(row.sign_in_time),
            _to_cairo(row.sign_out_time),
            f'{row.total_hours:.1f}' if row.total_hours is not None else '',
            _get_status_text(row.shift_status, row.status, lang=lang),
            yes_text if row.is_late else no_text,
            row.late_minutes or 0,
            day_off_text,
            absent_text,
            rate_text,
        ])

    # UTF-8 BOM so Excel opens Arabic CSV with correct encoding
    csv_bytes = '\ufeff' + buffer.getvalue()

    filename = f"{t('export_filename_csv', lang=lang)}.csv"
    encoded_filename = urllib.parse.quote(filename)

    return Response(
        csv_bytes.encode('utf-8'),
        mimetype='text/csv; charset=utf-8',
        headers={
            'Content-Disposition': f"attachment; filename=\"{filename}\"; filename*=UTF-8''{encoded_filename}"
        },
    )


def export_excel(rows, lang='ar', date_from=None, date_to=None):
    """Generates a formatted Excel (.xlsx) file with RTL layout for Arabic,
    LTR layout for English, custom headers, status badge highlighting, and column width auto-fitting."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = t('export_sheet_title', lang=lang)
    
    # RTL for Arabic, LTR for English
    is_arabic = (lang == 'ar')
    ws.views.sheetView[0].rightToLeft = is_arabic
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions
    font_title = Font(name="Segoe UI", size=15, bold=True, color="1E3A8A")
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_body = Font(name="Segoe UI", size=10, color="0F172A")
    
    fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_row_alt = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    fill_ok = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    font_ok = Font(name="Segoe UI", size=10, bold=True, color="166534")

    fill_warn = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
    font_warn = Font(name="Segoe UI", size=10, bold=True, color="854D0E")

    fill_info = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    font_info = Font(name="Segoe UI", size=10, bold=True, color="1E40AF")

    fill_danger = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    font_danger = Font(name="Segoe UI", size=10, bold=True, color="991B1B")

    align_center = Alignment(horizontal="center", vertical="center")
    align_side = Alignment(horizontal="right" if is_arabic else "left", vertical="center")

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # Sheet Title Banner
    ws.merge_cells('A1:L1')
    ws['A1'] = t('export_title_banner', lang=lang)
    ws['A1'].font = font_title
    ws['A1'].alignment = align_side
    ws.row_dimensions[1].height = 35

    headers = [
        t('export_col_date', lang=lang),
        t('export_col_employee', lang=lang),
        t('export_col_zone', lang=lang),
        t('export_col_signin', lang=lang),
        t('export_col_signout', lang=lang),
        t('export_col_total_hours', lang=lang),
        t('export_col_shift_status', lang=lang),
        t('export_col_late', lang=lang),
        t('export_col_late_minutes', lang=lang),
        t('export_col_extra_days', lang=lang),
        t('export_col_absent_days', lang=lang),
        t('export_col_attendance_rate', lang=lang),
    ]

    ws.append([])  # Blank row 2

    # Append Header Row at row 3
    ws.append(headers)
    ws.row_dimensions[3].height = 26
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=3, column=col_num)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border

    hours_suffix = t('export_hours_suffix', lang=lang)
    yes_text = t('export_yes', lang=lang)
    no_text = t('export_no', lang=lang)
    start_d = date_from or (min(r.date for r in rows) if rows else date.today())
    end_d = date_to or (max(r.date for r in rows) if rows else date.today())
    user_summaries = {}

    # Append Data Rows
    current_row = 4
    for row in rows:
        formatted_hours = f'{row.total_hours:.1f} {hours_suffix}' if row.total_hours is not None else '—'
        shift_status_text = _get_status_text(row.shift_status, row.status, lang=lang)
        late_text = yes_text if row.is_late else no_text

        if row.user_id not in user_summaries:
            user_summaries[row.user_id] = get_employee_attendance_summary(row.user, start_d, end_d)
        summary = user_summaries[row.user_id]

        is_day_off = not row.user.is_workday(row.date)
        day_off_text = yes_text if is_day_off else no_text

        absent_count = len(summary['absent_days'])
        absent_dates_str = ', '.join(d.strftime('%Y-%m-%d') for d in summary['absent_days'])
        absent_text = f"{absent_count} ({absent_dates_str})" if absent_count > 0 else t('no_absent_days', lang=lang)
        rate_text = f"{summary['attendance_rate']}%"

        data = [
            row.date.isoformat(),
            _sanitize_cell(row.user.name),
            _sanitize_cell(row.zone.name if row.zone else ''),
            _to_cairo(row.sign_in_time),
            _to_cairo(row.sign_out_time) if row.sign_out_time else '—',
            formatted_hours,
            shift_status_text,
            late_text,
            row.late_minutes or 0,
            day_off_text,
            absent_text,
            rate_text,
        ]
        ws.append(data)
        ws.row_dimensions[current_row].height = 22

        for col_num in range(1, len(data) + 1):
            cell = ws.cell(row=current_row, column=col_num)
            cell.font = font_body
            cell.alignment = align_center
            cell.border = thin_border
            if current_row % 2 == 1:
                cell.fill = fill_row_alt

            # Status Column (Column 7) styling
            if col_num == 7:
                if row.shift_status == 'full_shift':
                    cell.fill = fill_ok
                    cell.font = font_ok
                elif row.shift_status == 'short_shift':
                    cell.fill = fill_warn
                    cell.font = font_warn
                elif row.status == 'active':
                    cell.fill = fill_info
                    cell.font = font_info

            # Late Column (Column 8) styling
            if col_num == 8 and row.is_late:
                cell.fill = fill_danger
                cell.font = font_danger

            # Day off worked (Column 10) styling
            if col_num == 10 and is_day_off:
                cell.fill = fill_info
                cell.font = font_info

            # Absent count (Column 11) styling
            if col_num == 11 and absent_count > 0:
                cell.fill = fill_warn
                cell.font = font_warn

        current_row += 1

    # Auto-adjust column widths cleanly
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 6, 15)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"{t('export_filename_excel', lang=lang)}.xlsx"
    encoded_filename = urllib.parse.quote(filename)

    return Response(
        buffer.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={
            'Content-Disposition': f"attachment; filename=\"{filename}\"; filename*=UTF-8''{encoded_filename}"
        },
    )
